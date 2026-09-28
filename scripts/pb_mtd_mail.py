"""PB MTD Dashboard -> daily mail pack.

Builds the "PB MTD DASHBOARD" (key metrics, plant-wise by grade, retail
zone-wise & project) from one input workbook and writes everything needed for
the daily mail:

  * PB_MTD_Dashboard_<dd-Mon-yyyy>.xlsx  formatted attachment (values, no formulas,
                                         so phone / Outlook previews show numbers)
  * PB_MTD_Mail_<dd-Mon-yyyy>.html       Outlook-safe body (tables + inline CSS)
  * PB_MTD_Mail_<dd-Mon-yyyy>.eml        unsent draft: body + xlsx attached;
                                         double-click -> Outlook -> Send
  * PB_MTD_Dashboard_<dd-Mon-yyyy>.png   optional snapshot (--png, needs Playwright)

Usage:
  python scripts/pb_mtd_mail.py template <input.xlsx>        # blank input workbook
  python scripts/pb_mtd_mail.py build <input.xlsx> --out <dir>
        [--to a@x.com,b@y.com] [--cc ...] [--sender ...] [--png]

Input workbook sheets (see `template`):
  Meta  : key/value rows - as_on, latest_be (optional), dispatch_d1, btr,
          expected_closing_inv, doh, ageing, prev_month_invoiced
  Plant : Plant, Grade, BE, Orders, Invoiced, Conf Pending Invoice,
          Pending Orders, Physical Inv, DOH, Ageing, Critical Dia,
          PO Issued, PO Prod
  Zone  : Zone, Type, BE, Orders, Invoiced, Conf Pending Invoice, Pending Orders
          (Zone = "Project" for the project row)

Derived (never typed): KPI totals, Invoice % of BE, Production MTD (= sum of
PO Prod), PO Compliance, Order % vs BE, Retail+PTR subtotal, Grand totals.
Pending Orders is computed as Orders - Invoiced - Conf Pending when left blank.
If the zone rows don't add up to the plant grand total, an "Unmapped" row is
added so the zone table still ties to the headline numbers.
"""
from __future__ import annotations

import argparse
import html
import math
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from email.message import EmailMessage
from email.utils import formatdate
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ─── Column contracts ────────────────────────────────────────────────────────
PLANT_COLS = [
    "Plant", "Grade", "BE", "Orders", "Invoiced", "Conf Pending Invoice",
    "Pending Orders", "Physical Inv", "DOH", "Ageing", "Critical Dia",
    "PO Issued", "PO Prod",
]
ZONE_COLS = [
    "Zone", "Type", "BE", "Orders", "Invoiced", "Conf Pending Invoice",
    "Pending Orders",
]
QTY = ["BE", "Orders", "Invoiced", "Conf Pending Invoice", "Pending Orders"]
META_KEYS = {
    "as_on": "Report date (dd-mm-yyyy)",
    "latest_be": "Latest BE (MT) - blank = sum of plant BE",
    "dispatch_d1": "Dispatch D-1 (MT)",
    "btr": "Balance to produce - BTR (MT)",
    "expected_closing_inv": "Expected closing inventory (MT)",
    "doh": "Days of inventory - DOH (overall)",
    "ageing": "Ageing days (overall)",
    "prev_month_invoiced": "Prev month invoiced MTD (MT)",
}
# Pending check: |given - (Orders - Invoiced - Conf)| above this -> warning
PENDING_TOL_MT = 2.0
# Zone rows vs plant grand total: gap above this (any column) -> "Unmapped" row
ZONE_TIE_TOL_MT = 5.0

# ─── Traffic-light thresholds (upper bound inclusive -> fill) ────────────────
GREEN, AMBER, ORANGE, RED = "C6EFCE", "FFEB9C", "F8CBAD", "FFC7CE"
DOH_BANDS = [(7, GREEN), (20, AMBER), (30, ORANGE), (math.inf, RED)]
AGEING_BANDS = [(10, GREEN), (20, AMBER), (45, ORANGE), (math.inf, RED)]
# Order % vs BE: on-plan band green, overshoot amber/red, under-plan unfilled
ORDER_PCT_BANDS = [(0.90, None), (1.10, GREEN), (1.30, AMBER), (math.inf, RED)]

NAVY, HEAD_BG, TITLE_BG, TOTAL_BG = "002E5D", "DCE6F1", "FCE4D6", "FFC000"
KPI_GREEN = "00B050"


def band(value: float | None, bands: list[tuple[float, str | None]]) -> str | None:
    if value is None or pd.isna(value):
        return None
    for upper, colour in bands:
        if value <= upper:
            return colour
    return None


# ─── Model ───────────────────────────────────────────────────────────────────
@dataclass
class MtdReport:
    as_on: date
    meta: dict[str, float | None]
    plants: pd.DataFrame           # PLANT_COLS + "PO Compliance"
    plant_total: dict[str, float]
    zones: pd.DataFrame            # ZONE_COLS + "Order %", "_kind" (row|subtotal|total)
    kpi: list[tuple[str, str, str]]  # (label, value, note)
    warnings: list[str] = field(default_factory=list)

    @property
    def stamp(self) -> str:
        return self.as_on.strftime("%d-%b-%Y")

    @property
    def title(self) -> str:
        return (f"PB MTD DASHBOARD — {self.as_on.strftime('%b-%y').upper()} "
                f"(As on {self.stamp})")


def _num(v: Any) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str):
        v = v.replace(",", "").replace("%", "").strip()
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(f) else f


def _parse_date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    dt = pd.to_datetime(str(v).strip(), dayfirst=True, errors="coerce")
    if pd.isna(dt):
        raise ValueError(f"Meta.as_on is not a date: {v!r}")
    return dt.date()


def _frame(rows: list[list[Any]], cols: list[str], sheet: str) -> pd.DataFrame:
    if not rows:
        raise ValueError(f"Sheet '{sheet}' has no header row")
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    missing = [c for c in cols if c not in header]
    if missing:
        raise ValueError(f"Sheet '{sheet}' missing columns: {missing}")
    idx = [header.index(c) for c in cols]
    body = [[r[i] if i < len(r) else None for i in idx] for r in rows[1:]]
    body = [r for r in body if any(v not in (None, "") for v in r)]
    return pd.DataFrame(body, columns=cols)


def read_input(path: Path) -> tuple[date, dict[str, float | None], pd.DataFrame, pd.DataFrame]:
    wb = load_workbook(path, data_only=True)
    for s in ("Meta", "Plant", "Zone"):
        if s not in wb.sheetnames:
            raise ValueError(f"Input workbook needs a '{s}' sheet (has {wb.sheetnames})")
    meta_raw = {str(r[0]).strip(): r[1] for r in wb["Meta"].iter_rows(min_row=2, values_only=True)
                if r and r[0] is not None}
    if not meta_raw.get("as_on"):
        raise ValueError("Meta.as_on is required")
    as_on = _parse_date(meta_raw["as_on"])
    meta = {k: _num(meta_raw.get(k)) for k in META_KEYS if k != "as_on"}
    plants = _frame(list(wb["Plant"].iter_rows(values_only=True)), PLANT_COLS, "Plant")
    zones = _frame(list(wb["Zone"].iter_rows(values_only=True)), ZONE_COLS, "Zone")
    return as_on, meta, plants, zones


# ─── Derivations ─────────────────────────────────────────────────────────────
def _fmt_mt(v: float | None) -> str:
    return "" if v is None or pd.isna(v) else f"{v:,.0f}"


def _fmt_pct(v: float | None) -> str:
    return "" if v is None or pd.isna(v) else f"{v * 100:.0f}%"


def _fmt_pct1(v: float | None) -> str:
    return "" if v is None or pd.isna(v) else f"{v * 100:.1f}%"


def _fmt_days(v: float | None) -> str:
    if v is None or pd.isna(v):
        return ""
    return f"{v:.1f}" if v != round(v) else f"{v:.0f}"


def derive(as_on: date, meta: dict[str, float | None],
           plants: pd.DataFrame, zones: pd.DataFrame) -> MtdReport:
    warnings: list[str] = []
    p = plants.copy()
    p["Plant"] = p["Plant"].ffill().astype(str).str.strip()   # merged-cell style input
    p["Grade"] = p["Grade"].fillna("").astype(str).str.strip()
    p["Critical Dia"] = p["Critical Dia"].fillna("").astype(str).str.strip()
    num_cols = [c for c in PLANT_COLS if c not in ("Plant", "Grade", "Critical Dia")]
    for c in num_cols:
        p[c] = p[c].map(_num)

    calc = (p["Orders"].fillna(0) - p["Invoiced"].fillna(0)
            - p["Conf Pending Invoice"].fillna(0)).clip(lower=0)
    has_orders = p["Orders"].notna()
    for i in p.index[p["Pending Orders"].notna() & has_orders]:
        if abs(p.at[i, "Pending Orders"] - calc[i]) > PENDING_TOL_MT:
            warnings.append(
                f"Plant {p.at[i, 'Plant']} / {p.at[i, 'Grade']}: Pending Orders "
                f"{p.at[i, 'Pending Orders']:,.0f} ≠ Orders − Invoiced − Conf "
                f"({calc[i]:,.0f}) — check the row")
    fill = p["Pending Orders"].isna() & has_orders
    p.loc[fill, "Pending Orders"] = calc[fill]

    iss = p["PO Issued"]
    p["PO Compliance"] = (p["PO Prod"].fillna(0) / iss).where(iss.fillna(0) > 0, 0.0)

    tot = {c: float(p[c].fillna(0).sum()) for c in
           ["BE", "Orders", "Invoiced", "Conf Pending Invoice", "Pending Orders",
            "Physical Inv", "PO Issued", "PO Prod"]}
    tot["PO Compliance"] = tot["PO Prod"] / tot["PO Issued"] if tot["PO Issued"] else 0.0
    tot["DOH"] = meta.get("doh")
    tot["Ageing"] = meta.get("ageing")

    latest_be = meta.get("latest_be")
    if latest_be is None:
        latest_be = tot["BE"]
    elif abs(latest_be - tot["BE"]) > 0.5:
        warnings.append(f"Meta.latest_be {latest_be:,.0f} ≠ plant BE total {tot['BE']:,.0f}")

    # Zone table: rows -> Retail+PTR subtotal -> Project -> (Unmapped) -> Grand total
    z = zones.copy()
    z["Zone"] = z["Zone"].ffill().fillna("").astype(str).str.strip()
    z["Type"] = z["Type"].fillna("").astype(str).str.strip()
    for c in QTY:
        z[c] = z[c].map(_num)
    is_proj = z["Zone"].str.lower().eq("project")
    retail, proj = z[~is_proj].copy(), z[is_proj].copy()
    retail["_kind"], proj["_kind"] = "row", "row"
    sub = {"Zone": "Retail+PTR", "Type": "", "_kind": "subtotal",
           **{c: float(retail[c].fillna(0).sum()) for c in QTY}}
    parts = [retail, pd.DataFrame([sub]), proj]
    zsum = {c: sub[c] + float(proj[c].fillna(0).sum()) for c in QTY}
    gap = {c: tot[c] - zsum[c] for c in QTY}
    if any(abs(v) > ZONE_TIE_TOL_MT for v in gap.values()):
        shown = {c: (v if abs(v) > ZONE_TIE_TOL_MT else None) for c, v in gap.items()}
        parts.append(pd.DataFrame([{"Zone": "Unmapped", "Type": "", "_kind": "row", **shown}]))
        warnings.append(
            "Zone table doesn't tie to plant totals — added 'Unmapped' row: "
            + ", ".join(f"{c} {v:+,.0f}" for c, v in gap.items() if abs(v) > ZONE_TIE_TOL_MT))
    parts.append(pd.DataFrame([{"Zone": "GRAND TOTAL", "Type": "", "_kind": "total",
                                **{c: tot[c] for c in QTY}}]))
    zt = pd.concat(parts, ignore_index=True)
    be = zt["BE"]
    zt["Order %"] = (zt["Orders"] / be).where(be.fillna(0) > 0)

    prev = meta.get("prev_month_invoiced")
    inv_pct = tot["Invoiced"] / latest_be if latest_be else None
    kpi = [
        ("Latest BE (MT)", _fmt_mt(latest_be), ""),
        ("Total Orders Received (MT)", _fmt_mt(tot["Orders"]), ""),
        ("Invoiced MTD (MT)", _fmt_mt(tot["Invoiced"]),
         f"Prev Month Invoiced (MTD): {prev / 1000:.1f} KMT" if prev else ""),
        ("Dispatch D-1 (MT)", _fmt_mt(meta.get("dispatch_d1")), ""),
        ("Invoice % of BE", _fmt_pct(inv_pct), ""),
        ("Physical Inventory (MT)", _fmt_mt(tot["Physical Inv"]), ""),
        ("Production MTD (MT)", _fmt_mt(tot["PO Prod"]), ""),
        ("Balance to Produce (BTR) (MT)", _fmt_mt(meta.get("btr")), ""),
        ("Expected Closing Inv (MT)", _fmt_mt(meta.get("expected_closing_inv")), ""),
        ("Days of Inventory — DOH", _fmt_days(meta.get("doh")), ""),
        ("Ageing (Days)", _fmt_days(meta.get("ageing")), ""),
    ]
    for k in ("dispatch_d1", "btr", "expected_closing_inv", "doh", "ageing"):
        if meta.get(k) is None:
            warnings.append(f"Meta.{k} is blank — KPI tile will be empty")
    return MtdReport(as_on, meta, p, tot, zt, kpi, warnings)


# ─── Excel attachment ────────────────────────────────────────────────────────
_thin = Side(style="thin", color="A6A6A6")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _fill(hex_: str | None) -> PatternFill:
    return PatternFill("solid", fgColor=hex_) if hex_ else PatternFill(fill_type=None)


def _put(ws, r: int, c: int, v: Any, *, fmt: str | None = None, bold: bool = False,
         bg: str | None = None, align: Alignment = RIGHT, color: str = "000000",
         size: int = 10) -> None:
    cell = ws.cell(row=r, column=c)
    cell.value = None if v is None or (isinstance(v, float) and pd.isna(v)) else v
    cell.font = Font(name="Calibri", size=size, bold=bold, color=color)
    cell.alignment = align
    cell.border = BORDER
    if bg:
        cell.fill = _fill(bg)
    if fmt:
        cell.number_format = fmt


def _band_row(ws, r: int, c1: int, c2: int, text: str) -> None:
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    _put(ws, r, c1, text, bold=True, bg=TITLE_BG, align=CENTER, size=11)
    for c in range(c1 + 1, c2 + 1):
        ws.cell(row=r, column=c).border = BORDER


MT, PCT0, PCT1 = "#,##0", "0%", "0.0%"


def build_xlsx(rep: MtdReport, out: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PB MTD"
    ws.sheet_view.showGridLines = False
    widths = [16, 11, 10, 10, 10, 11, 13, 11, 8, 9, 34, 2, 11, 11, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 1
    _band_row(ws, r, 1, 15, rep.title)
    ws.row_dimensions[r].height = 22
    ws.cell(row=r, column=1).font = Font(name="Calibri", size=14, bold=True, color=NAVY)

    # ① KPIs — 11 tiles across A..K
    r = 3
    _band_row(ws, r, 1, 11, "① KEY METRICS — MTD SNAPSHOT")
    raw = {
        0: (rep.meta.get("latest_be") or rep.plant_total["BE"], MT),
        1: (rep.plant_total["Orders"], MT), 2: (rep.plant_total["Invoiced"], MT),
        3: (rep.meta.get("dispatch_d1"), MT),
        4: (rep.plant_total["Invoiced"] / (rep.meta.get("latest_be") or rep.plant_total["BE"] or 1), PCT0),
        5: (rep.plant_total["Physical Inv"], MT), 6: (rep.plant_total["PO Prod"], MT),
        7: (rep.meta.get("btr"), MT), 8: (rep.meta.get("expected_closing_inv"), MT),
        9: (rep.meta.get("doh"), "0"), 10: (rep.meta.get("ageing"), "0.0"),
    }
    for i, (label, _, note) in enumerate(rep.kpi):
        c = i + 1
        _put(ws, r + 1, c, label, bold=True, bg=HEAD_BG, align=CENTER, size=8)
        v, fmt = raw[i]
        _put(ws, r + 2, c, v, fmt=fmt, bold=True, align=CENTER, color=KPI_GREEN, size=16)
        _put(ws, r + 3, c, note or None, align=CENTER, size=7, color="C00000")
    ws.row_dimensions[r + 1].height = 34
    ws.row_dimensions[r + 2].height = 30
    ws.row_dimensions[r + 3].height = 22

    # ② Plant-wise + PO block (M..O)
    r = 8
    _band_row(ws, r, 1, 11, "② PLANT-WISE PERFORMANCE (MT) — BY GRADE")
    _band_row(ws, r, 13, 15, "PO STATUS (MT)")
    heads = ["Plant", "Grade", "BE", "Orders", "Invoiced", "Conf. Pending Invoice",
             "Pending Orders (Non-Conf+SFDC)", "Physical Inv", "DOH", "Ageing (Days)",
             "Critical Dia"]
    for c, h in enumerate(heads, 1):
        _put(ws, r + 1, c, h, bold=True, bg=HEAD_BG, align=CENTER)
    for c, h in zip((13, 14, 15), ("PO Issued", "PO Prod", "PO Compliance")):
        _put(ws, r + 1, c, h, bold=True, bg=HEAD_BG, align=CENTER)
    ws.row_dimensions[r + 1].height = 30

    r0 = r + 2
    for i, (_, d) in enumerate(rep.plants.iterrows()):
        rr = r0 + i
        vals = [d["Plant"], d["Grade"], d["BE"], d["Orders"], d["Invoiced"],
                d["Conf Pending Invoice"], d["Pending Orders"], d["Physical Inv"],
                d["DOH"], d["Ageing"], d["Critical Dia"] or None]
        for c, v in enumerate(vals, 1):
            bg = (band(v, DOH_BANDS) if c == 9 else
                  band(v, AGEING_BANDS) if c == 10 else None)
            _put(ws, rr, c, v, fmt=MT if 3 <= c <= 10 else None, bg=bg,
                 align=LEFT if c in (1, 2, 11) else RIGHT,
                 bold=c == 1, size=8 if c == 11 else 10,
                 color="C00000" if c == 11 else "000000")
        _put(ws, rr, 13, d["PO Issued"], fmt=MT)
        _put(ws, rr, 14, d["PO Prod"], fmt=MT)
        _put(ws, rr, 15, d["PO Compliance"], fmt=PCT1)
    # merge Plant cells for consecutive rows of the same plant
    names = list(rep.plants["Plant"])
    start = 0
    for i in range(1, len(names) + 1):
        if i == len(names) or names[i] != names[start]:
            if i - start > 1:
                ws.merge_cells(start_row=r0 + start, start_column=1,
                               end_row=r0 + i - 1, end_column=1)
            start = i
    rt = r0 + len(names)
    t = rep.plant_total
    tot_vals = ["GRAND TOTAL", None, t["BE"], t["Orders"], t["Invoiced"],
                t["Conf Pending Invoice"], t["Pending Orders"], t["Physical Inv"],
                t["DOH"], t["Ageing"], None]
    for c, v in enumerate(tot_vals, 1):
        _put(ws, rt, c, v, fmt=MT if 3 <= c <= 10 else None, bold=True, bg=TOTAL_BG,
             align=LEFT if c <= 2 else RIGHT)
    _put(ws, rt, 13, t["PO Issued"], fmt=MT, bold=True, bg=TOTAL_BG)
    _put(ws, rt, 14, t["PO Prod"], fmt=MT, bold=True, bg=TOTAL_BG)
    _put(ws, rt, 15, t["PO Compliance"], fmt=PCT0, bold=True, bg=TOTAL_BG)

    # ③ Zone-wise
    r = rt + 2
    _band_row(ws, r, 1, 8, "③ RETAIL (ZONE-WISE) & PROJECT PERFORMANCE (MT)")
    zh = ["Zone", "Type", "BE", "Orders", "Order % vs BE", "Invoiced",
          "Conf. Pending Invoice", "Pending Orders (Non-Conf+SFDC)"]
    for c, h in enumerate(zh, 1):
        _put(ws, r + 1, c, h, bold=True, bg=HEAD_BG, align=CENTER)
    ws.row_dimensions[r + 1].height = 30
    z0 = r + 2
    zones = rep.zones
    for i, d in zones.iterrows():
        rr = z0 + i
        strong = d["_kind"] != "row"
        bg = TOTAL_BG if d["_kind"] == "total" else ("FFF2CC" if strong else None)
        vals = [d["Zone"], d["Type"] or None, d["BE"], d["Orders"], d["Order %"],
                d["Invoiced"], d["Conf Pending Invoice"], d["Pending Orders"]]
        for c, v in enumerate(vals, 1):
            cbg = bg
            if c == 5 and d["_kind"] != "total":
                cbg = band(v, ORDER_PCT_BANDS) or bg
            _put(ws, rr, c, v, fmt=PCT0 if c == 5 else (MT if c >= 3 else None),
                 bold=strong or c == 1, bg=cbg, align=LEFT if c <= 2 else RIGHT)
    # merge Zone cells only across consecutive plain rows of the same zone
    zlabels = [f"{z}|{k}" if k == "row" else f"{z}|{k}|{j}"
               for j, (z, k) in enumerate(zip(zones["Zone"], zones["_kind"]))]
    start = 0
    for i in range(1, len(zlabels) + 1):
        if i == len(zlabels) or zlabels[i] != zlabels[start]:
            if i - start > 1:
                ws.merge_cells(start_row=z0 + start, start_column=1,
                               end_row=z0 + i - 1, end_column=1)
            start = i

    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ─── HTML mail body (Outlook-safe: tables, inline styles, bgcolor) ───────────
_TD = "border:1px solid #BFBFBF;padding:3px 6px;font:12px Calibri,Arial,sans-serif;"
_TH = _TD + "background:#DCE6F1;font-weight:bold;text-align:center;"


def _td(v: str, *, bg: str | None = None, bold: bool = False, align: str = "right",
        extra: str = "", attrs: str = "") -> str:
    style = _TD + f"text-align:{align};" + ("font-weight:bold;" if bold else "") + extra
    bga = f' bgcolor="#{bg}"' if bg else ""
    if bg:
        style += f"background:#{bg};"
    return f'<td{bga}{attrs} style="{style}">{v}</td>'


def _section(title: str, colspan: int) -> str:
    return (f'<tr><td colspan="{colspan}" bgcolor="#{TITLE_BG}" style="{_TD}'
            f'background:#{TITLE_BG};font-weight:bold;text-align:center;font-size:13px;">'
            f'{html.escape(title)}</td></tr>')


def build_html(rep: MtdReport) -> str:
    e = html.escape
    t = rep.plant_total
    out: list[str] = [
        '<!DOCTYPE html><html><head><meta charset="utf-8">'
        f"<title>{e(rep.title)}</title></head>"
        '<body style="margin:0;padding:12px;background:#FFFFFF;">',
        f'<p style="font:13px Calibri,Arial,sans-serif;margin:0 0 10px 0;">Dear All,<br><br>'
        f"Please find below the PB MTD dashboard as on <b>{e(rep.stamp)}</b> "
        f"(Excel attached).</p>",
        f'<table cellpadding="0" cellspacing="0" style="border-collapse:collapse;">'
        f'<tr><td style="font:bold 16px Calibri,Arial,sans-serif;color:#{NAVY};'
        f'padding:4px 0 8px 0;">{e(rep.title)}</td></tr></table>',
    ]

    # ① KPI tiles: two rows (6 + 5) so it fits a 760px mail pane
    out.append('<table cellpadding="0" cellspacing="0" style="border-collapse:collapse;'
               'margin-bottom:12px;">' + _section("① KEY METRICS — MTD SNAPSHOT", 6))
    for chunk in (rep.kpi[:6], rep.kpi[6:]):
        out.append("<tr>")
        for label, value, note in chunk:
            note_html = (f'<br><span style="font-size:10px;color:#C00000;font-weight:normal;">'
                         f"{e(note)}</span>" if note else "")
            out.append(
                f'<td width="125" style="{_TD}text-align:center;vertical-align:top;'
                f'padding:6px 4px;"><span style="font-size:10px;font-weight:bold;'
                f'color:#404040;">{e(label)}</span><br>'
                f'<span style="font-size:22px;font-weight:bold;color:#{KPI_GREEN};">'
                f"{e(value) or '—'}</span>{note_html}</td>")
        out.append(f'<td style="{_TD}"></td>' * (6 - len(chunk)) + "</tr>")
    out.append("</table>")

    # ② Plant-wise
    heads = ["Plant", "Grade", "BE", "Orders", "Invoiced", "Conf. Pending Invoice",
             "Pending Orders (Non-Conf+SFDC)", "Physical Inv", "DOH", "Ageing (Days)",
             "Critical Dia", "PO Issued", "PO Prod", "PO Compliance"]
    out.append('<table cellpadding="0" cellspacing="0" style="border-collapse:collapse;'
               'margin-bottom:12px;">'
               + _section("② PLANT-WISE PERFORMANCE (MT) — BY GRADE", len(heads))
               + "<tr>" + "".join(f'<td style="{_TH}">{e(h)}</td>' for h in heads) + "</tr>")
    p = rep.plants
    spans = p.groupby("Plant", sort=False).size().to_dict()
    seen: set[str] = set()
    for _, row in p.iterrows():
        out.append("<tr>")
        if row["Plant"] not in seen:
            seen.add(row["Plant"])
            out.append(_td(e(row["Plant"]), bold=True, align="left",
                           attrs=f' rowspan="{spans[row["Plant"]]}"', extra="vertical-align:middle;"))
        out.append(_td(e(row["Grade"]), align="left"))
        for c in ("BE", "Orders", "Invoiced", "Conf Pending Invoice", "Pending Orders",
                  "Physical Inv"):
            out.append(_td(_fmt_mt(row[c])))
        out.append(_td(_fmt_mt(row["DOH"]), bg=band(row["DOH"], DOH_BANDS)))
        out.append(_td(_fmt_mt(row["Ageing"]), bg=band(row["Ageing"], AGEING_BANDS)))
        out.append(_td(e(row["Critical Dia"]), align="left",
                       extra="font-size:10px;color:#C00000;max-width:220px;"))
        out.append(_td(_fmt_mt(row["PO Issued"])) + _td(_fmt_mt(row["PO Prod"]))
                   + _td(_fmt_pct1(row["PO Compliance"])))
        out.append("</tr>")
    tv = ["GRAND TOTAL", "", _fmt_mt(t["BE"]), _fmt_mt(t["Orders"]), _fmt_mt(t["Invoiced"]),
          _fmt_mt(t["Conf Pending Invoice"]), _fmt_mt(t["Pending Orders"]),
          _fmt_mt(t["Physical Inv"]), _fmt_mt(t["DOH"]), _fmt_mt(t["Ageing"]), "",
          _fmt_mt(t["PO Issued"]), _fmt_mt(t["PO Prod"]), _fmt_pct(t["PO Compliance"])]
    out.append("<tr>" + "".join(_td(e(v), bg=TOTAL_BG, bold=True,
                                    align="left" if i < 2 else "right")
                                for i, v in enumerate(tv)) + "</tr></table>")

    # ③ Zone-wise
    zh = ["Zone", "Type", "BE", "Orders", "Order % vs BE", "Invoiced",
          "Conf. Pending Invoice", "Pending Orders (Non-Conf+SFDC)"]
    out.append('<table cellpadding="0" cellspacing="0" style="border-collapse:collapse;'
               'margin-bottom:12px;">'
               + _section("③ RETAIL (ZONE-WISE) & PROJECT PERFORMANCE (MT)", len(zh))
               + "<tr>" + "".join(f'<td style="{_TH}">{e(h)}</td>' for h in zh) + "</tr>")
    z = rep.zones
    for i, row in z.iterrows():
        kind = row["_kind"]
        strong = kind != "row"
        bg = TOTAL_BG if kind == "total" else ("FFF2CC" if kind == "subtotal" else None)
        out.append("<tr>")
        first_of_zone = kind != "row" or i == 0 or z.at[i - 1, "Zone"] != row["Zone"] \
            or z.at[i - 1, "_kind"] != "row"
        if first_of_zone:
            n = 1
            while kind == "row" and i + n < len(z) and z.at[i + n, "Zone"] == row["Zone"] \
                    and z.at[i + n, "_kind"] == "row":
                n += 1
            out.append(_td(e(row["Zone"]), bold=True, bg=bg, align="left",
                           attrs=f' rowspan="{n}"' if n > 1 else "",
                           extra="vertical-align:middle;"))
        out.append(_td(e(row["Type"]), bg=bg, bold=strong, align="left"))
        out.append(_td(_fmt_mt(row["BE"]), bg=bg, bold=strong))
        out.append(_td(_fmt_mt(row["Orders"]), bg=bg, bold=strong))
        pbg = band(row["Order %"], ORDER_PCT_BANDS) if kind != "total" else None
        out.append(_td(_fmt_pct(row["Order %"]), bg=pbg or bg, bold=strong))
        for c in ("Invoiced", "Conf Pending Invoice", "Pending Orders"):
            out.append(_td(_fmt_mt(row[c]), bg=bg, bold=strong))
        out.append("</tr>")
    out.append("</table>")
    out.append('<p style="font:10px Calibri,Arial,sans-serif;color:#7F7F7F;margin:4px 0 12px 0;">'
               "All quantities in MT. DOH: ≤7 green · 8–20 amber · 21–30 orange · &gt;30 red. "
               "Ageing: ≤10 green · 11–20 amber · 21–45 orange · &gt;45 red. "
               "Order % vs BE: 90–110% green · 110–130% amber · &gt;130% red.</p>"
               '<p style="font:13px Calibri,Arial,sans-serif;margin:0;">Regards,</p>'
               "</body></html>")
    return "".join(out)


# ─── .eml draft ──────────────────────────────────────────────────────────────
def build_eml(rep: MtdReport, body_html: str, xlsx: Path, out: Path, *,
              to: str = "", cc: str = "", sender: str = "") -> Path:
    msg = EmailMessage()
    msg["Subject"] = f"PB MTD Dashboard — {rep.as_on.strftime('%b-%y')} (As on {rep.stamp})"
    if sender:
        msg["From"] = sender
    if to:
        msg["To"] = to
    if cc:
        msg["Cc"] = cc
    msg["Date"] = formatdate(localtime=True)
    msg["X-Unsent"] = "1"          # Outlook opens it as an editable draft
    msg.set_content(f"PB MTD Dashboard as on {rep.stamp}. Excel attached.")
    msg.add_alternative(body_html, subtype="html")
    msg.add_attachment(xlsx.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       filename=xlsx.name)
    out.write_bytes(bytes(msg))
    return out


def render_png(html_path: Path, png: Path) -> Path:
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        try:
            b = pw.chromium.launch()
        except Exception:  # pinned Playwright without its bundled browser
            exe = next(Path("/opt/pw-browsers").glob("chromium-*/chrome-linux*/chrome"), None)
            if exe is None:
                raise
            b = pw.chromium.launch(executable_path=str(exe))
        page = b.new_page(viewport={"width": 1400, "height": 900}, device_scale_factor=2)
        page.goto(html_path.resolve().as_uri())
        page.screenshot(path=str(png), full_page=True)
        b.close()
    return png


# ─── Input template ──────────────────────────────────────────────────────────
def write_template(out: Path, *, as_on: str = "", meta: dict[str, Any] | None = None,
                   plants: list[list[Any]] | None = None,
                   zones: list[list[Any]] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Meta"
    ws.append(["key", "value", "description"])
    meta = meta or {}
    for k, desc in META_KEYS.items():
        ws.append([k, as_on if k == "as_on" else meta.get(k), desc])
    for name, cols, rows in (("Plant", PLANT_COLS, plants), ("Zone", ZONE_COLS, zones)):
        s = wb.create_sheet(name)
        s.append(cols)
        for r in rows or []:
            s.append(r)
    for s in wb.worksheets:
        for c in s[1]:
            c.font = Font(bold=True)
            c.fill = _fill(HEAD_BG)
        for i in range(1, s.max_column + 1):
            s.column_dimensions[get_column_letter(i)].width = 16
    wb["Meta"].column_dimensions["C"].width = 44
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def build_pack(inp: Path, out_dir: Path, *, to: str = "", cc: str = "", sender: str = "",
               png: bool = False) -> tuple[MtdReport, dict[str, Path]]:
    rep = derive(*read_input(inp))
    out_dir.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}
    files["xlsx"] = build_xlsx(rep, out_dir / f"PB_MTD_Dashboard_{rep.stamp}.xlsx")
    body = build_html(rep)
    files["html"] = out_dir / f"PB_MTD_Mail_{rep.stamp}.html"
    files["html"].write_text(body, encoding="utf-8")
    files["eml"] = build_eml(rep, body, files["xlsx"], out_dir / f"PB_MTD_Mail_{rep.stamp}.eml",
                             to=to, cc=cc, sender=sender)
    if png:
        files["png"] = render_png(files["html"], out_dir / f"PB_MTD_Dashboard_{rep.stamp}.png")
    return rep, files


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    t = sub.add_parser("template", help="write a blank input workbook")
    t.add_argument("path", type=Path)
    b = sub.add_parser("build", help="build xlsx + html + eml from an input workbook")
    b.add_argument("input", type=Path)
    b.add_argument("--out", type=Path, default=Path(".workspace/daily_mail"))
    b.add_argument("--to", default="")
    b.add_argument("--cc", default="")
    b.add_argument("--sender", default="")
    b.add_argument("--png", action="store_true", help="also render a PNG snapshot")
    a = ap.parse_args(argv)
    if a.cmd == "template":
        print(write_template(a.path))
        return 0
    rep, files = build_pack(a.input, a.out, to=a.to, cc=a.cc, sender=a.sender, png=a.png)
    for w in rep.warnings:
        print(f"WARNING: {w}", file=sys.stderr)
    for k, pth in files.items():
        print(f"{k}: {pth}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
